#!/usr/bin/env python3
"""Post-process the Word file for the Journal of Population Economics (Springer).
Times New Roman 11 pt, 1.5 line spacing, 1-inch margins, automatic page numbers (footer PAGE field),
single-spaced tables, blank author metadata (anonymized copy)."""
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

path = sys.argv[1]
d = Document(path)
for sec in d.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    for tag in ('w:headerReference', 'w:footerReference'):
        for e in sec._sectPr.findall(qn(tag)): sec._sectPr.remove(e)
    ft = sec.footer
    ft.is_linked_to_previous = False
    p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
    for r in list(p.runs): r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; run._r.append(it)
# fonts and spacing on all styles that carry text
for st in d.styles:
    try:
        if st.type != 1: continue
        f = st.font
        f.name = 'Times New Roman'
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), 'Times New Roman')
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'): rf.attrib.pop(qn(a), None)
        f.color.rgb = None
        if st.name.startswith('Heading') or st.name in ('Title', 'Abstract Title'):
            f.bold = True; f.italic = (st.name == 'Heading 2')
            f.size = Pt({'Title': 14, 'Heading 1': 12, 'Heading 2': 11, 'Heading 3': 11}.get(st.name, 11))
        else:
            f.size = Pt(11)
        if st.name in ('Normal', 'Body Text', 'First Paragraph', 'Compact', 'Abstract', 'Block Text'):
            st.paragraph_format.line_spacing = 1.5
    except Exception:
        pass
cp = d.core_properties
cp.author = ''; cp.last_modified_by = ''; cp.title = "Grandmothers' care and mothers' time: coresident grandmothers in India"; cp.comments = ''; cp.keywords = ''
d.save(path)

# Tables: single-spaced, 9 pt, autofit to text width
d = Document(path)
for t in d.tables:
    tblPr = t._tbl.tblPr
    for tag in ('w:tblW', 'w:tblLayout'):
        for e in tblPr.findall(qn(tag)): tblPr.remove(e)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000'); tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'autofit'); tblPr.append(lay)
    grid = t._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc in grid.findall(qn('w:gridCol')): gc.attrib.pop(qn('w:w'), None)
    for row in t.rows:
        for c in row.cells:
            tcPr = c._tc.get_or_add_tcPr()
            for e in tcPr.findall(qn('w:tcW')): tcPr.remove(e)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                for r in p.runs: r.font.size = Pt(9)
d.save(path)

# Number the display equation: centered equation, right-aligned "(n)" (Word has no automatic number for pandoc's equation)
from docx.enum.text import WD_TAB_ALIGNMENT
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
d = Document(path)
n_eq = 0
for p in d.paragraphs:
    mp = p._p.findall('{%s}oMathPara' % M)
    if len(mp) == 1 and not p.text.strip():
        n_eq += 1
        om = mp[0].find('{%s}oMath' % M)
        p._p.remove(mp[0])
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t')
        p._p.append(om)
        p.add_run('\t(%d)' % n_eq)
d.save(path)
print('numbered %d display equations' % n_eq)

# Number the table captions ("Table 1: ...") so they match the in-text citations (shared tool, idempotent)
import os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'tools'))
from number_captions_docx import number_captions
print('numbered %d table captions' % number_captions(path))
