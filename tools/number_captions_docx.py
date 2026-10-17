#!/usr/bin/env python3
"""Number the table captions of the Word file ("Table 1: ...") so they match the in-text citations.
Usage: python3 tools/number_captions_docx.py file.docx   (idempotent; called from tools/fix_docx.py)"""
import sys
from docx import Document

def number_captions(path):
    d = Document(path)
    k = 0
    for para in d.paragraphs:
        if para.style.name == 'Table Caption' and not para.text.startswith('Table '):
            k += 1
            if para.runs:
                para.runs[0].text = 'Table %d: ' % k + para.runs[0].text
            else:
                para.add_run('Table %d: ' % k)
    d.save(path)
    return k

if __name__ == '__main__':
    print('numbered %d captions' % number_captions(sys.argv[1]))
